#!/usr/bin/env python3

from __future__ import annotations

import argparse
import json
import logging
import re
from io import BytesIO
from pathlib import Path
from typing import Any

import pandas as pd
import requests
from openpyxl import load_workbook

API_BASE = "https://idg.ine.gob.bo/api"
DEFAULT_OUTPUT_DIR = Path("temporal/ficha_prueba")
DEFAULT_CAMPOS = Path("recursos/campos.json")

HEADERS = {
    "User-Agent": "Mozilla/5.0 (X11; Linux x86_64; rv:144.0) Gecko/20100101 Firefox/144.0",
    "Accept": "application/json, text/plain, */*",
    "Accept-Language": "en-US,en;q=0.5",
    "Content-Type": "application/json",
    "Origin": "https://idg.ine.gob.bo",
    "Referer": "https://idg.ine.gob.bo/geoportal",
}

LOGGER = logging.getLogger("descargar_ficha_prueba")


def slugify(value: str) -> str:
    value = value.strip().lower()
    value = value.replace("%", "porcentaje")
    value = re.sub(r"[^\w\s/-]", "", value)
    value = value.replace("/", "_")
    value = re.sub(r"[\s-]+", "_", value)
    return value.strip("_")


class TokenManager:
    def __init__(self, session: requests.Session, timeout: float):
        self.session = session
        self.timeout = timeout
        self._token: str | None = None

    def get_token(self, force_refresh: bool = False) -> str:
        if self._token is None or force_refresh:
            response = self.session.post(
                f"{API_BASE}/auth/acceso",
                headers=HEADERS,
                json={},
                timeout=self.timeout,
            )
            response.raise_for_status()
            token = response.json().get("token")
            if not token:
                raise RuntimeError("auth/acceso did not return token")
            self._token = token
        return self._token


def build_headers(token: str) -> dict[str, str]:
    return {**HEADERS, "Authorization": f"Bearer {token}"}


def request_json(
    session: requests.Session,
    token_manager: TokenManager,
    path: str,
    payload: dict[str, Any],
    timeout: float,
) -> dict[str, Any]:
    token = token_manager.get_token()
    response = session.post(
        f"{API_BASE}{path}",
        headers=build_headers(token),
        json=payload,
        timeout=timeout,
    )
    if response.status_code in (401, 403):
        token = token_manager.get_token(force_refresh=True)
        response = session.post(
            f"{API_BASE}{path}",
            headers=build_headers(token),
            json=payload,
            timeout=timeout,
        )
    response.raise_for_status()
    return response.json()


def request_excel(
    session: requests.Session,
    token_manager: TokenManager,
    payload: dict[str, Any],
    timeout: float,
) -> bytes:
    token = token_manager.get_token()
    response = session.post(
        f"{API_BASE}/ficha/generar-excel",
        headers=build_headers(token),
        json=payload,
        timeout=timeout,
    )
    if response.status_code in (401, 403):
        token = token_manager.get_token(force_refresh=True)
        response = session.post(
            f"{API_BASE}/ficha/generar-excel",
            headers=build_headers(token),
            json=payload,
            timeout=timeout,
        )
    response.raise_for_status()
    return response.content


def cell_value_to_json(value: Any) -> Any:
    if hasattr(value, "item"):
        return value.item()
    return value


def workbook_cells_dict(workbook_bytes: bytes) -> dict[str, dict[str, Any]]:
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)
    result: dict[str, dict[str, Any]] = {}
    for worksheet in workbook.worksheets:
        cells: dict[str, Any] = {}
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value not in (None, ""):
                    cells[cell.coordinate] = cell_value_to_json(cell.value)
        result[worksheet.title] = cells
    return result


def parse_field_value(worksheet, cell_ref: str) -> Any:
    value = worksheet[cell_ref].value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return cell_value_to_json(value)

    row = worksheet[cell_ref].row
    col = worksheet[cell_ref].column
    for next_col in range(col + 1, worksheet.max_column + 1):
        candidate = worksheet.cell(row=row, column=next_col).value
        if isinstance(candidate, (int, float)) and not isinstance(candidate, bool):
            return cell_value_to_json(candidate)

    return cell_value_to_json(value)


def parse_known_fields(workbook_bytes: bytes, campos: dict[str, str]) -> dict[str, Any]:
    worksheet = load_workbook(
        BytesIO(workbook_bytes), data_only=True, read_only=True
    ).active
    return {
        field_name: parse_field_value(worksheet, cell_ref)
        for cell_ref, field_name in campos.items()
    }


def infer_simple_tables(workbook_bytes: bytes) -> dict[str, Any]:
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)
    inferred: dict[str, Any] = {}

    for worksheet in workbook.worksheets:
        rows = list(worksheet.iter_rows(values_only=True))
        sheet_key = slugify(worksheet.title)
        sheet_data: dict[str, Any] = {}

        for row in rows:
            values = list(row)
            text_positions = [
                idx
                for idx, value in enumerate(values)
                if isinstance(value, str) and value.strip()
            ]
            numeric_positions = [
                idx
                for idx, value in enumerate(values)
                if isinstance(value, (int, float)) and not isinstance(value, bool)
            ]
            if len(text_positions) == 1 and len(numeric_positions) == 1:
                label = values[text_positions[0]]
                number = values[numeric_positions[0]]
                key = slugify(label)
                if key:
                    sheet_data[key] = cell_value_to_json(number)

        if sheet_data:
            inferred[sheet_key] = sheet_data

    return inferred


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Descarga y parsea fichas de un manzano de prueba."
    )
    parser.add_argument("codigo", help="Código de manzano, por ejemplo 00629873876-A.")
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
        help="Directorio donde guardar JSON, parquet y xlsx de prueba.",
    )
    parser.add_argument(
        "--campos",
        type=Path,
        default=DEFAULT_CAMPOS,
        help="JSON con mappings de celdas conocidas para base y vivienda.",
    )
    parser.add_argument(
        "--timeout",
        type=float,
        default=120.0,
        help="Timeout por request en segundos.",
    )
    parser.add_argument(
        "--log-level",
        default="INFO",
        choices=["DEBUG", "INFO", "WARNING", "ERROR"],
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    logging.basicConfig(
        level=getattr(logging, args.log_level),
        format="%(asctime)s %(levelname)s %(message)s",
        force=True,
    )

    args.output_dir.mkdir(parents=True, exist_ok=True)
    with open(args.campos, "r", encoding="utf-8") as f:
        campos = json.load(f)

    if not isinstance(campos, dict) or "base" not in campos or "vivienda" not in campos:
        raise RuntimeError("campos.json debe tener llaves 'base' y 'vivienda'")

    session = requests.Session()
    token_manager = TokenManager(session=session, timeout=args.timeout)

    LOGGER.info("Validando manzano %s", args.codigo)
    verification = request_json(
        session,
        token_manager,
        "/ficha/verificarValidar",
        {"mara": 2024, "codigos": [args.codigo]},
        args.timeout,
    )

    LOGGER.info("Descargando ficha base")
    base_bytes = request_excel(
        session,
        token_manager,
        {"mara": "2024", "codigos": [args.codigo], "vivienda": False},
        args.timeout,
    )

    LOGGER.info("Descargando ficha vivienda")
    vivienda_bytes = request_excel(
        session,
        token_manager,
        {"mara": "2024", "codigos": [args.codigo], "vivienda": True},
        args.timeout,
    )

    base_xlsx = args.output_dir / f"{args.codigo}_base.xlsx"
    vivienda_xlsx = args.output_dir / f"{args.codigo}_vivienda.xlsx"
    base_xlsx.write_bytes(base_bytes)
    vivienda_xlsx.write_bytes(vivienda_bytes)

    base_known_fields = parse_known_fields(base_bytes, campos["base"])
    vivienda_known_fields = parse_known_fields(vivienda_bytes, campos["vivienda"])
    ficha_unificada = {
        "codigo": args.codigo,
        **base_known_fields,
        **vivienda_known_fields,
    }
    base_cells = workbook_cells_dict(base_bytes)
    vivienda_cells = workbook_cells_dict(vivienda_bytes)
    vivienda_inferred = infer_simple_tables(vivienda_bytes)

    summary_row = {
        "codigo": args.codigo,
        "validado": bool(verification.get("validado", False)),
        "personas": verification.get("cantidad_personas"),
        "viviendas": verification.get("cantidad_viviendas"),
        "mensaje": verification.get("mensaje"),
    }

    pd.DataFrame([summary_row]).to_parquet(
        args.output_dir / f"{args.codigo}_poblacion.parquet"
    )
    pd.DataFrame([ficha_unificada]).to_parquet(
        args.output_dir / f"{args.codigo}_fichas.parquet"
    )

    result = {
        "codigo": args.codigo,
        "verificacion": verification,
        "poblacion_vivienda": summary_row,
        "ficha_unificada": ficha_unificada,
        "ficha_base": {
            "archivo": str(base_xlsx),
            "campos_conocidos": base_known_fields,
            "celdas_no_vacias": base_cells,
        },
        "ficha_vivienda": {
            "archivo": str(vivienda_xlsx),
            "campos_conocidos": vivienda_known_fields,
            "campos_inferidos": vivienda_inferred,
            "celdas_no_vacias": vivienda_cells,
        },
    }

    output_json = args.output_dir / f"{args.codigo}_fichas.json"
    with open(output_json, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    LOGGER.info("Listo. Archivos generados en %s", args.output_dir)
    LOGGER.info(
        "Resumen: validado=%s personas=%s viviendas=%s",
        summary_row["validado"],
        summary_row["personas"],
        summary_row["viviendas"],
    )
    LOGGER.info("JSON consolidado: %s", output_json)


if __name__ == "__main__":
    main()
